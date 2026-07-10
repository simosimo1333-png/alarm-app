# -*- coding: utf-8 -*-
"""YAONA！今後の運営方針案 PowerPoint 生成スクリプト"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- カラーパレット（ネイビー・ブルー・ライトグレー基調）----
NAVY = RGBColor(0x1B, 0x2A, 0x4C)
BLUE = RGBColor(0x2E, 0x5F, 0xA3)
LIGHT_BLUE = RGBColor(0xDC, 0xE8, 0xF7)
MID_BLUE = RGBColor(0x5B, 0x87, 0xC5)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
GRAY_LINE = RGBColor(0xD5, 0xDB, 0xE3)
GRAY_HEADER = RGBColor(0xC9, 0xD1, 0xDC)
TEXT_DARK = RGBColor(0x33, 0x3A, 0x45)
TEXT_GRAY = RGBColor(0x6B, 0x74, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Yu Gothic"

SW, SH = 13.333, 7.5
MX = 0.55            # 左右マージン
CW = SW - 2 * MX     # コンテンツ幅

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _set_font(run, size, bold, color, name=FONT, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def add_text(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of dicts {text, size, bold, color, align, space_before, line_spacing}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", PP_ALIGN.LEFT)
        if ln.get("space_before"):
            p.space_before = Pt(ln["space_before"])
        p.line_spacing = ln.get("line_spacing", 1.15)
        runs = ln.get("runs")
        if runs:
            for rt in runs:
                r = p.add_run()
                r.text = rt["text"]
                _set_font(r, rt.get("size", ln.get("size", 12)),
                          rt.get("bold", ln.get("bold", False)),
                          rt.get("color", ln.get("color", TEXT_DARK)))
        else:
            r = p.add_run()
            r.text = ln["text"]
            _set_font(r, ln.get("size", 12), ln.get("bold", False),
                      ln.get("color", TEXT_DARK))
    return box


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None,
              line_w=0.75, radius=None):
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    if radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return sp


def shape_text(sp, lines):
    tf = sp.text_frame
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", PP_ALIGN.CENTER)
        if ln.get("space_before"):
            p.space_before = Pt(ln["space_before"])
        p.line_spacing = ln.get("line_spacing", 1.1)
        r = p.add_run()
        r.text = ln["text"]
        _set_font(r, ln.get("size", 12), ln.get("bold", False),
                  ln.get("color", TEXT_DARK))


PAGE_NO = [0]


def content_slide(title, lead=None):
    """共通ヘッダー・フッター付きスライド"""
    PAGE_NO[0] += 1
    slide = prs.slides.add_slide(BLANK)
    # 背景
    bg = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    bg.line.fill.background()
    # タイトル
    add_shape(slide, MSO_SHAPE.RECTANGLE, MX, 0.42, 0.09, 0.52, fill=NAVY)
    add_text(slide, MX + 0.22, 0.40, CW - 0.3, 0.6,
             [{"text": title, "size": 21, "bold": True, "color": NAVY}])
    add_shape(slide, MSO_SHAPE.RECTANGLE, MX, 1.05, CW, 0.018, fill=GRAY_LINE)
    if lead:
        add_text(slide, MX, 1.22, CW, 0.4,
                 [{"text": lead, "size": 12.5, "color": TEXT_GRAY}])
    # フッター
    add_text(slide, MX, 7.12, 6.0, 0.3,
             [{"text": "YAONA！ 今後の運営方針案（庁内検討用）",
               "size": 9, "color": TEXT_GRAY}])
    add_text(slide, SW - 1.1, 7.12, 0.6, 0.3,
             [{"text": str(PAGE_NO[0]), "size": 10, "color": TEXT_GRAY,
               "align": PP_ALIGN.RIGHT}])
    return slide


def bullet_lines(items, size=12.5, color=TEXT_DARK, sp=8):
    return [{"text": "・" + t, "size": size, "color": color,
             "space_before": 0 if i == 0 else sp, "line_spacing": 1.2}
            for i, t in enumerate(items)]


# =====================================================================
# 1. 表紙
# =====================================================================
PAGE_NO[0] += 1
s = prs.slides.add_slide(BLANK)
add_shape(s, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
# 上部ラベル
chip = add_shape(s, MSO_SHAPE.RECTANGLE, MX, 0.55, 1.75, 0.42,
                 fill=None, line=NAVY, line_w=1.0)
shape_text(chip, [{"text": "庁内検討用資料", "size": 11, "bold": True,
                   "color": NAVY}])
# 中央ネイビー帯
band = add_shape(s, MSO_SHAPE.RECTANGLE, 0, 2.55, SW, 2.15, fill=NAVY)
add_text(s, 0, 2.95, SW, 1.0,
         [{"text": "YAONA！ 今後の運営方針案", "size": 34, "bold": True,
           "color": WHITE, "align": PP_ALIGN.CENTER}])
add_shape(s, MSO_SHAPE.RECTANGLE, SW / 2 - 0.6, 3.78, 1.2, 0.025,
          fill=MID_BLUE)
add_text(s, 0, 3.95, SW, 0.5,
         [{"text": "参加者の取扱 ・ 参加後のつながり ・ サポーター制度の整理",
           "size": 15, "color": LIGHT_BLUE, "align": PP_ALIGN.CENTER}])
add_text(s, 0, 5.35, SW, 0.4,
         [{"text": "2026年7月", "size": 13, "color": TEXT_GRAY,
           "align": PP_ALIGN.CENTER}])

# =====================================================================
# 2. 検討事項の全体像
# =====================================================================
s = content_slide("検討事項の全体像",
                  "YAONA！の持続的な運営に向けて、次の3点を整理する。")
col_w = (CW - 0.7) / 3
items2 = [
    ("01", "参加者の取扱",
     ["参加は「1人2回まで」", "新しい若手に開かれた\n入口であり続ける"]),
    ("02", "参加後のつながり",
     ["つながりの昇華・\nスピンアウトを促す", "市は“軽い後押し”に徹する"]),
    ("03", "サポーター制度",
     ["役割と謝礼の明確化", "謝礼に加え“学びの機会”\nとしての価値を整理"]),
]
for i, (no, ttl, lines) in enumerate(items2):
    x = MX + i * (col_w + 0.35)
    card = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.95, col_w, 4.0,
                     fill=LIGHT_GRAY, line=GRAY_LINE, radius=0.05)
    add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.12, 1.95, col_w - 0.24, 0.07,
              fill=NAVY)
    circ = add_shape(s, MSO_SHAPE.OVAL, x + col_w / 2 - 0.36, 2.35,
                     0.72, 0.72, fill=BLUE)
    shape_text(circ, [{"text": no, "size": 18, "bold": True, "color": WHITE}])
    add_text(s, x + 0.2, 3.3, col_w - 0.4, 0.5,
             [{"text": ttl, "size": 16.5, "bold": True, "color": NAVY,
               "align": PP_ALIGN.CENTER}])
    body = []
    for j, t in enumerate(lines):
        for k, seg in enumerate(t.split("\n")):
            body.append({"text": seg, "size": 12, "color": TEXT_DARK,
                         "align": PP_ALIGN.CENTER,
                         "space_before": (12 if k == 0 and j > 0 else 0),
                         "line_spacing": 1.2})
    add_text(s, x + 0.2, 4.0, col_w - 0.4, 1.8, body)

# =====================================================================
# 3〜5. 現状スライド（良い点／懸念／必要な対応）
# =====================================================================
def status_slide(title, lead, good, concern, action):
    s = content_slide(title, lead)
    col_w = (CW - 0.7) / 3
    cols = [
        ("◎ 良い点", BLUE, WHITE, good),
        ("△ 懸念", GRAY_HEADER, NAVY, concern),
        ("→ 必要な対応", NAVY, WHITE, action),
    ]
    for i, (hd, hfill, htxt, lines) in enumerate(cols):
        x = MX + i * (col_w + 0.35)
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.85, col_w, 4.2,
                  fill=LIGHT_GRAY if i != 2 else LIGHT_BLUE,
                  line=GRAY_LINE, radius=0.045)
        head = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.85, col_w, 0.62,
                         fill=hfill, radius=0.045)
        shape_text(head, [{"text": hd, "size": 15, "bold": True,
                           "color": htxt}])
        add_text(s, x + 0.28, 2.75, col_w - 0.56, 3.1,
                 bullet_lines(lines, size=12.5, sp=12))
    return s


status_slide(
    "現状① 参加者の取扱",
    "リピーターの良さと、新規参加者への影響を整理する。",
    ["リピーターはグループワーク\nを円滑にしてくれる".replace("\n", ""),
     "場の雰囲気づくりに貢献している"],
    ["既存の輪が強くなると、新規参加者が入りにくくなる",
     "“常連の場”になると、入口としての機能が弱まる"],
    ["YAONA！を新しい若手に開き続ける",
     "参加ルールを明確にし、運営で場をコントロールする"],
)

status_slide(
    "現状② 参加後のつながり",
    "セミナー後の展開をどう促すかが課題となっている。",
    ["セミナー内で良いつながりが生まれている",
     "「もっと話したい」という声もある"],
    ["その後の展開は参加者任せで、見えにくい",
     "つながりが一過性で終わる可能性がある"],
    ["市が主導するのではなく、自発的な動きを促す仕組みをつくる",
     "“軽い後押し”のメニューを用意する"],
)

status_slide(
    "現状③ サポーター制度",
    "有志の善意に依存しない、持続可能な制度への整理が必要。",
    ["過去登壇者等が有志・ボランティア的に協力してくれている",
     "現場をよく知る心強い存在"],
    ["本業がある中での協力に依存する構造は望ましくない",
     "役割や位置づけが不明確なまま負担が偏るおそれ"],
    ["役割・謝礼を明確にする",
     "サポーター自身の“学びの価値”もあわせて整理する"],
)

# =====================================================================
# 6. 基本方針
# =====================================================================
s = content_slide("基本方針")
msg = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 1.4, CW, 1.15,
                fill=NAVY, radius=0.08)
shape_text(msg, [
    {"text": "YAONA！は「新しい若手が参加する入口」。",
     "size": 16.5, "bold": True, "color": WHITE},
    {"text": "そこで生まれたつながりは、参加者自身の意思で次に広がる場とする。",
     "size": 16.5, "bold": True, "color": WHITE, "space_before": 4},
])
# フロー図
steps = [
    ("YAONA！", "出会い・学びの入口", NAVY, WHITE, WHITE),
    ("自発的なつながり", "参加者の意思で継続", BLUE, WHITE, WHITE),
    ("小さな活動", "座談会・勉強会など", MID_BLUE, WHITE, WHITE),
    ("地域・職場への還元", "気づきと行動の広がり", LIGHT_BLUE, NAVY, NAVY),
]
bw, bh, gap = 2.62, 1.5, 0.45
total = 4 * bw + 3 * gap
x = MX + (CW - total) / 2
y = 3.15
for i, (ttl, cap, fill, tc, cc) in enumerate(steps):
    box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh,
                    fill=fill, radius=0.09)
    shape_text(box, [
        {"text": ttl, "size": 14.5, "bold": True, "color": tc},
        {"text": cap, "size": 10.5, "color": cc, "space_before": 4},
    ])
    if i < 3:
        ar = add_shape(s, MSO_SHAPE.RIGHT_ARROW, x + bw + 0.05,
                       y + bh / 2 - 0.16, 0.35, 0.32, fill=GRAY_HEADER)
    x += bw + gap
# 注記
note = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + 1.5, 5.25,
                 CW - 3.0, 0.75, fill=LIGHT_GRAY, line=GRAY_LINE,
                 radius=0.12)
shape_text(note, [{"text": "市の役割は「主導」ではなく、参加者が動き出すための“軽い後押し”",
                   "size": 13.5, "bold": True, "color": NAVY}])

# =====================================================================
# 7. 参加者の取扱方針
# =====================================================================
s = content_slide("参加者の取扱方針",
                  "新しい若手が入りやすい場を保つため、参加ルールを明確にする。")
cards7 = [
    ("1", "参加は1人2回まで", "参加上限を募集要項に明記する"),
    ("2", "初参加者を優先", "申込多数の場合は初参加者を優先する"),
    ("3", "2回目参加者の役割", "グループ内で場を温める役割を期待する"),
    ("4", "配置の工夫", "既存のつながりが固まりすぎないよう\nグループ編成を調整する"),
]
cw2 = (CW - 0.4) / 2
for i, (no, ttl, sub) in enumerate(cards7):
    x = MX + (i % 2) * (cw2 + 0.4)
    y = 1.75 + (i // 2) * 1.75
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw2, 1.5,
              fill=LIGHT_GRAY, line=GRAY_LINE, radius=0.07)
    circ = add_shape(s, MSO_SHAPE.OVAL, x + 0.28, y + 0.42, 0.62, 0.62,
                     fill=BLUE)
    shape_text(circ, [{"text": no, "size": 17, "bold": True, "color": WHITE}])
    lines = [{"text": ttl, "size": 15.5, "bold": True, "color": NAVY}]
    for k, seg in enumerate(sub.split("\n")):
        lines.append({"text": seg, "size": 11.5, "color": TEXT_GRAY,
                      "space_before": 5 if k == 0 else 0,
                      "line_spacing": 1.15})
    add_text(s, x + 1.15, y + 0.28, cw2 - 1.4, 1.1, lines)
note = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.5, CW, 0.8,
                 fill=LIGHT_BLUE, radius=0.1)
shape_text(note, [{"text": "上限設定は「排除」ではなく、新規参加者が入りやすい場を保つための運営上の工夫",
                   "size": 13.5, "bold": True, "color": NAVY}])

# =====================================================================
# 8〜10. 参加後支援案 A/B/C
# =====================================================================
def support_slide(label, title, lead, content_builder, roles, point):
    s = content_slide(title, lead)
    chip = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 1.62, 2.0, 0.44,
                     fill=BLUE, radius=0.5)
    shape_text(chip, [{"text": label, "size": 11.5, "bold": True,
                       "color": WHITE}])
    # 内容カード（左）
    lx, ly, lw, lh = MX, 2.2, 7.45, 3.05
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, lx, ly, lw, lh,
              fill=LIGHT_GRAY, line=GRAY_LINE, radius=0.045)
    add_text(s, lx + 0.3, ly + 0.22, 1.5, 0.4,
             [{"text": "内容", "size": 14, "bold": True, "color": BLUE}])
    content_builder(s, lx, ly, lw, lh)
    # 市の役割カード（右）
    rx, rw = lx + lw + 0.35, CW - lw - 0.35
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, rx, ly, rw, lh,
              fill=WHITE, line=GRAY_LINE, radius=0.045)
    head = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, rx, ly, rw, 0.6,
                     fill=NAVY, radius=0.045)
    shape_text(head, [{"text": "市の役割（軽い後押し）", "size": 13.5,
                       "bold": True, "color": WHITE}])
    add_text(s, rx + 0.28, ly + 0.85, rw - 0.56, lh - 1.0,
             bullet_lines(roles, size=12.5, sp=12))
    # ポイント
    note = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 5.55, CW, 0.75,
                     fill=LIGHT_BLUE, radius=0.1)
    shape_text(note, [{
        "text": "ポイント：" + point, "size": 13, "bold": True,
        "color": NAVY}])
    return s


def content_a(s, lx, ly, lw, lh):
    add_text(s, lx + 0.3, ly + 0.62, lw - 0.6, 1.1, bullet_lines([
        "終了時に「もっと話したいテーマ」「一緒に考えたいこと」をカードに記入",
        "近い関心を持つ参加者同士が、自然につながれるようにする",
    ], size=12.5, sp=8))
    # ミニフロー
    fsteps = ["カードに記入", "テーマを見える化", "関心の近い人がつながる"]
    fw, fh, fg = 2.05, 0.62, 0.32
    fx = lx + 0.3
    fy = ly + lh - 0.95
    for i, t in enumerate(fsteps):
        box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, fx, fy, fw, fh,
                        fill=WHITE, line=BLUE, line_w=1.0, radius=0.16)
        shape_text(box, [{"text": t, "size": 11, "bold": True,
                          "color": NAVY}])
        if i < 2:
            add_shape(s, MSO_SHAPE.RIGHT_ARROW, fx + fw + 0.04,
                      fy + fh / 2 - 0.11, 0.24, 0.22, fill=GRAY_HEADER)
        fx += fw + fg


support_slide(
    "参加後支援 案A", "参加後支援案A：関心テーマの見える化",
    "「話したいこと」を起点に、参加者同士が自然につながるきっかけをつくる。",
    content_a,
    ["カード様式の準備", "テーマの整理", "希望者同士の接点づくり"],
    "つなぐのは市ではなく“テーマ”。市は道具と機会を用意する。",
)


def content_b(s, lx, ly, lw, lh):
    add_text(s, lx + 0.3, ly + 0.62, lw - 0.6, 0.8, bullet_lines([
        "参加者が“小さく”始められる活動の形を用意する",
        "ハードルの低い形から、無理なく一歩を踏み出せるようにする",
    ], size=12.5, sp=8))
    chips = ["3人程度の座談会", "職場見学", "勉強会"]
    fw, fh, fg = 2.1, 0.62, 0.35
    fx = lx + 0.3
    fy = ly + lh - 0.95
    for t in chips:
        box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, fx, fy, fw, fh,
                        fill=WHITE, line=BLUE, line_w=1.0, radius=0.16)
        shape_text(box, [{"text": t, "size": 11.5, "bold": True,
                          "color": NAVY}])
        fx += fw + fg


support_slide(
    "参加後支援 案B", "参加後支援案B：小さな自主企画の芽出し",
    "参加者発の活動を、小さな単位から始められるように支える。",
    content_b,
    ["企画テンプレートの提供", "相談窓口の設置", "会場情報の提供"],
    "市は主催者にはならず、参加者発の活動を後押しする。",
)


def content_c(s, lx, ly, lw, lh):
    add_text(s, lx + 0.3, ly + 0.62, lw - 0.6, 0.8, bullet_lines([
        "参加者発で生まれた小さな動きを、次回YAONA！の冒頭や休憩時間に紹介",
    ], size=12.5))
    add_text(s, lx + 0.3, ly + 1.35, 1.5, 0.4,
             [{"text": "効果", "size": 14, "bold": True, "color": BLUE}])
    add_text(s, lx + 0.3, ly + 1.75, lw - 0.6, 1.1, bullet_lines([
        "参加者の自発的な活動が“見える”ようになる",
        "次の参加者にも波及し、動き出す人が増える",
    ], size=12.5, sp=8))


support_slide(
    "参加後支援 案C", "参加後支援案C：次回YAONA！での活動紹介枠",
    "参加者発の動きを可視化し、次の一歩を後押しする。",
    content_c,
    ["紹介機会の提供に留める", "（企画・運営は参加者が担う）"],
    "動きの共有が、次に踏み出す人の心理的ハードルを下げる。",
)

# =====================================================================
# 11. サポーター制度の方針
# =====================================================================
s = content_slide("サポーター制度の方針")
msg = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 1.35, CW, 0.85,
                fill=NAVY, radius=0.1)
shape_text(msg, [{"text": "正式に役割を依頼する場合は謝礼を支払い、「謝礼」と「学び」を別の価値として整理する",
                  "size": 15, "bold": True, "color": WHITE}])
cw11 = (CW - 0.5) / 2
cards11 = [
    ("謝礼 ― 役割への対価", BLUE, [
        "正式に役割を依頼する場合は謝礼を支払う",
        "役割と責任の所在を明確にする",
    ]),
    ("学び ― 参加そのものの価値", NAVY, [
        "若手社会人の“生の声”を聞くことができる",
        "支援を通じて、自身の学び・研修の機会になる",
        "若手との対話を通じ、地域や職場の実情を知る機会になる",
    ]),
]
for i, (hd, hfill, lines) in enumerate(cards11):
    x = MX + i * (cw11 + 0.5)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, cw11, 3.15,
              fill=LIGHT_GRAY, line=GRAY_LINE, radius=0.05)
    head = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, cw11, 0.62,
                     fill=hfill, radius=0.05)
    shape_text(head, [{"text": hd, "size": 14.5, "bold": True,
                       "color": WHITE}])
    add_text(s, x + 0.32, 3.45, cw11 - 0.64, 2.1,
             bullet_lines(lines, size=12.5, sp=12))
plus = add_shape(s, MSO_SHAPE.OVAL, MX + cw11 + 0.02, 3.85, 0.46, 0.46,
                 fill=WHITE, line=NAVY, line_w=1.25)
shape_text(plus, [{"text": "＋", "size": 16, "bold": True, "color": NAVY}])
add_text(s, MX, 5.95, CW, 0.4,
         [{"text": "依頼の際は、謝礼とあわせて“学びの機会”としての価値も伝える",
           "size": 12.5, "color": TEXT_GRAY, "align": PP_ALIGN.CENTER}])

# =====================================================================
# 12. 清原氏によるサポーター向け講座案・次回試行事項
# =====================================================================
s = content_slide("サポーター向け講座案と次回試行事項")
add_text(s, MX, 1.25, 3.0, 0.4,
         [{"text": "■ 清原氏による事前講座", "size": 14, "bold": True,
           "color": NAVY}])
# テーマチップ
themes = ["コーチング", "傾聴", "対話の場づくり"]
tx = MX + 3.1
for t in themes:
    chip = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx, 1.24, 1.85, 0.42,
                     fill=LIGHT_BLUE, radius=0.5)
    shape_text(chip, [{"text": t, "size": 11.5, "bold": True,
                       "color": NAVY}])
    tx += 2.05
# 3ステップフロー
fsteps = [
    ("STEP 1", "事前講座", "清原氏による講義"),
    ("STEP 2", "YAONA！当日", "学びを現場で実践"),
    ("STEP 3", "終了後", "振り返り・共有"),
]
fw, fh, fg = 3.55, 1.25, 0.5
total = 3 * fw + 2 * fg
fx = MX + (CW - total) / 2
fy = 1.95
fills = [NAVY, BLUE, MID_BLUE]
for i, (stp, ttl, cap) in enumerate(fsteps):
    box = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, fx, fy, fw, fh,
                    fill=fills[i], radius=0.08)
    shape_text(box, [
        {"text": stp, "size": 10, "bold": True, "color": LIGHT_BLUE},
        {"text": ttl, "size": 15, "bold": True, "color": WHITE,
         "space_before": 2},
        {"text": cap, "size": 10.5, "color": WHITE, "space_before": 2},
    ])
    if i < 2:
        add_shape(s, MSO_SHAPE.RIGHT_ARROW, fx + fw + 0.08,
                  fy + fh / 2 - 0.15, 0.34, 0.3, fill=GRAY_HEADER)
    fx += fw + fg
# 区切り
add_shape(s, MSO_SHAPE.RECTANGLE, MX, 3.55, CW, 0.015, fill=GRAY_LINE)
add_text(s, MX, 3.72, 6.0, 0.4,
         [{"text": "■ 次回までに決めること", "size": 14, "bold": True,
           "color": NAVY}])
todos = [
    "参加上限の募集要項への明記",
    "2回目参加者の配置（グループ編成）ルール",
    "参加後支援案（A〜C）の試行内容",
    "サポーターの役割・謝礼・講座時間",
]
cw12 = (CW - 0.4) / 2
for i, t in enumerate(todos):
    x = MX + (i % 2) * (cw12 + 0.4)
    y = 4.25 + (i // 2) * 0.95
    card = add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw12, 0.78,
                     fill=LIGHT_GRAY, line=GRAY_LINE, radius=0.12)
    chk = add_shape(s, MSO_SHAPE.OVAL, x + 0.22, y + 0.19, 0.4, 0.4,
                    fill=BLUE)
    shape_text(chk, [{"text": "✓", "size": 13, "bold": True,
                      "color": WHITE}])
    add_text(s, x + 0.8, y + 0.22, cw12 - 1.0, 0.5,
             [{"text": t, "size": 13, "bold": True, "color": TEXT_DARK}])

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "docs", "YAONA_運営方針案.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("saved:", out)
